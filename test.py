import os
os.environ['KMP_DUPLICATE_LIB_OK']='True' # Add this line AT THE VERY TOP

import streamlit as st
import uuid
import re
import os
from openai import OpenAI, APIError, AuthenticationError, RateLimitError
import pandas as pd
from datetime import datetime
from langdetect import detect, LangDetectException # As per document section 3.3
import numpy as np

# --- New Imports for File Processing ---
from io import BytesIO

# --- PyPDF2 for PDF processing (Document 2.1.1, 3.3) ---
py_pdf_available = False
try:
    import PyPDF2
    py_pdf_available = True
except ImportError:
    PyPDF2 = None

# --- python-docx for DOCX processing (Document 2.1.1, 3.3) ---
python_docx_available = False
try:
    import docx
    python_docx_available = True
except ImportError:
    docx = None

# --- BeautifulSoup for HTML processing (Document 2.1.1, 3.3) ---
beautiful_soup_available = False
try:
    from bs4 import BeautifulSoup
    beautiful_soup_available = True
except ImportError:
    BeautifulSoup = None

# --- Pillow for image manipulation (needed by EasyOCR) ---
pillow_available = False
try:
    from PIL import Image
    pillow_available = True
except ImportError:
    Image = None

# --- EasyOCR for OCR (Document 2.1.1 notes Tesseract, code uses EasyOCR) ---
easy_ocr_available = False
easyocr_reader_en_fr = None # Global reader instance for caching
try:
    import easyocr
    easy_ocr_available = True
    if easy_ocr_available:
        try:
            # Initialize for English and French as common languages
            easyocr_reader_en_fr = easyocr.Reader(['en', 'fr'], gpu=False)
        except Exception as e:
            # st.warning(f"EasyOCR Reader init failed: {e}") # Suppress for cleaner UI startup
            easy_ocr_available = False
            easyocr_reader_en_fr = None
except ImportError:
    easyocr = None

# --- Requests for Web Scraping ---
requests_available = False
try:
    import requests
    requests_available = True
except ImportError:
    requests = None

# --- FAISS for Vector Storage (Document 2.1.1, 3.3) ---
faiss_available = False
try:
    import faiss
    faiss_available = True
except ImportError:
    faiss = None

# Call set_page_config() as the FIRST Streamlit command
st.set_page_config(layout="wide")

# --- Global constants ---
EMBEDDING_DIMS = {
    "text-embedding-3-small": 1536,
    "text-embedding-ada-002": 1536,
    "text-embedding-3-large": 3072,
}

# --- Initialize session state ---
if 'documents' not in st.session_state:
    st.session_state.documents = {}
if 'doc_counter' not in st.session_state:
    st.session_state.doc_counter = 0
if 'openai_api_key' not in st.session_state:
    st.session_state.openai_api_key = None
if 'embedding_model' not in st.session_state:
    st.session_state.embedding_model = "text-embedding-3-small" # Default model as per doc 2.1.1
if 'rag_enabled' not in st.session_state:
    st.session_state.rag_enabled = True

if 'faiss_index' not in st.session_state:
    st.session_state.faiss_index = None
if 'faiss_index_dim' not in st.session_state:
    st.session_state.faiss_index_dim = None
if 'chunk_store' not in st.session_state:
    st.session_state.chunk_store = []

if 'file_uploader_key_suffix' not in st.session_state:
    st.session_state.file_uploader_key_suffix = 0
if 'web_url_input_key_suffix' not in st.session_state:
    st.session_state.web_url_input_key_suffix = 0
if 'doc_name_manual_input_key_suffix' not in st.session_state:
    st.session_state.doc_name_manual_input_key_suffix = 0
if 'doc_text_area_input_key_suffix' not in st.session_state:
    st.session_state.doc_text_area_input_key_suffix = 0

# Conceptual Guardrails State (Section 3.4)
if 'guardrails_enabled' not in st.session_state:
    st.session_state.guardrails_enabled = {
        "conli": False, "cove": False, "contextcheck": False,
        "detectpii": False, "unusual_prompt": False
    }

# --- Helper Functions ---
# (query_preprocessor.py equivalent from doc 3.1)
def preprocess_text(text):
    text = re.sub(r'\s+', ' ', text)
    text = text.strip()
    text = re.sub(r'<[^>]+>', '', text) # Basic HTML tag stripping
    return text

def simple_chunker(text, chunk_size=200, overlap=50):
    words = text.split()
    chunks = []
    current_pos = 0
    while current_pos < len(words):
        end_pos = min(current_pos + chunk_size, len(words))
        chunk_words = words[current_pos:end_pos]
        chunks.append(" ".join(chunk_words))
        if end_pos == len(words):
            break
        current_pos += (chunk_size - overlap)
        if current_pos >= len(words): # Ensure we don't go past the end
            break
    return [c for c in chunks if c.strip()]

# (embedding_generator.py equivalent from doc 3.1)
def get_embedding(text, api_key, model="text-embedding-3-small"):
    if not api_key:
        st.error("OpenAI API Key not provided for embedding generation.")
        return None
    try:
        client = OpenAI(api_key=api_key)
        text = text.replace("\n", " ")
        response = client.embeddings.create(input=[text], model=model)
        return response.data[0].embedding
    except AuthenticationError:
        st.error("OpenAI API Key is invalid. Could not generate embeddings.")
        return None
    except RateLimitError:
        st.error("Rate limit exceeded for OpenAI Embeddings API.")
        return None
    except APIError as e:
        st.error(f"OpenAI API error during embedding: {e}")
        return None
    except Exception as e:
        st.error(f"An unexpected error occurred during embedding: {str(e)}")
        return None

# (chunk_retriever.py equivalent from doc 3.1)
def faiss_similarity_search(query_embedding, faiss_index, chunk_store_list, top_k=3, similarity_threshold=0.7):
    if query_embedding is None or faiss_index is None or not chunk_store_list or faiss_index.ntotal == 0:
        return []
    
    query_embedding_np = np.array([query_embedding], dtype='float32')
    scores, indices = faiss_index.search(query_embedding_np, top_k)
    
    results = []
    if indices.size > 0:
        for i in range(indices.shape[1]): 
            faiss_id = indices[0, i]
            score = float(scores[0, i])
            
            if faiss_id == -1: continue

            if score >= similarity_threshold:
                if 0 <= faiss_id < len(chunk_store_list):
                    chunk_data = chunk_store_list[faiss_id].copy() # Return a copy
                    chunk_data['score'] = score 
                    results.append(chunk_data)
                else:
                    st.warning(f"FAISS returned ID {faiss_id} which is out of bounds for chunk_store (size {len(chunk_store_list)}). Skipping.")
    return results

# --- Conceptual Guardrail Validator Functions (Section 3.4) ---
# These are placeholders. Real implementation would involve specific libraries/logic.
def conli_guard_validator(response_text, context_chunks):
    # Validates coherence phrase by phrase against verified knowledge (context)
    # st.info("[Guardrail - conli-guard]: Checking response-context coherence...")
    return True, "conli-guard: OK" 

def cove_guard_validator(response_text, context_chunks, query):
    # Verifies if the response can be logically justified by a verification chain from context
    # st.info("[Guardrail - cove-guard]: Checking logical justification...")
    return True, "cove-guard: OK"

def contextcheck_guard_validator(response_text, conversation_history):
    # Ensures response considers conversational context (e.g., history, previous intentions)
    # st.info("[Guardrail - contextcheck-guard]: Checking conversational context consistency...")
    return True, "contextcheck-guard: OK"

def detectpii_guard_validator(text_to_check):
    # Detects PII
    # st.info(f"[Guardrail - detectpii-guard]: Scanning for PII in '{text_to_check[:30]}...'")
    if "john.doe@email.com" in text_to_check.lower(): # Dummy PII check
         return False, "detectpii-guard: Potential PII detected (dummy check)."
    return True, "detectpii-guard: OK"

def unusual_prompt_guard_validator(user_query):
    # Detects suspicious or tricky prompts
    # st.info("[Guardrail - unusual-prompt-guard]: Analyzing prompt for suspicious patterns...")
    if "ignore previous instructions" in user_query.lower(): # Dummy check
        return False, "unusual-prompt-guard: Potentially problematic prompt detected."
    return True, "unusual-prompt-guard: OK"

# (Prompt Constructor + LLM call - implicit from doc 3.1, pipeline step 7)
def real_llm_generation_with_openai(query, context_chunks, api_key, model="gpt-4o"): # Model updated as per doc
    if not api_key:
        return "OpenAI API Key not provided. Please enter it in the sidebar."

    # --- Conceptual Guardrail Integration Point (Input) ---
    if st.session_state.guardrails_enabled.get("unusual_prompt", False):
        is_prompt_ok, prompt_guard_msg = unusual_prompt_guard_validator(query)
        st.sidebar.caption(f"Unusual Prompt Guard: {prompt_guard_msg}")
        if not is_prompt_ok:
            return f"Query blocked by unusual-prompt-guard: {prompt_guard_msg}"
    
    if st.session_state.guardrails_enabled.get("detectpii", False):
        is_pii_ok_query, pii_guard_msg_query = detectpii_guard_validator(query)
        st.sidebar.caption(f"PII Guard (Query): {pii_guard_msg_query}")
        # Could decide to block or warn based on policy

    try:
        client = OpenAI(api_key=api_key)
        # Prompt Construction (doc 3.1 prompt_constructor.py)
        if not context_chunks or not st.session_state.rag_enabled:
            prompt_messages = [
                {"role": "system", "content": "You are a helpful assistant."},
                {"role": "user", "content": f"Answer the following question: {query}"}
            ]
        else:
            context_text = "\n\n---\n\n".join([chunk['text'] for chunk in context_chunks])
            prompt_messages = [
                {"role": "system", "content": "You are a helpful assistant. Answer the user's question based ONLY on the provided context. If the answer is not in the context, say you cannot answer based on the provided information. Cite the source document ID or name if possible from the context."},
                {"role": "user", "content": f"Here is the context from one or more documents:\n\n---\n{context_text}\n---\n\nBased on this context, answer the question: {query}"}
            ]
        
        response = client.chat.completions.create(
            model=model, # Using model parameter, can be gpt-4, gpt-4o as per doc
            messages=prompt_messages,
            temperature=0.2, # Low temperature for factual RAG
            max_tokens=350
        )
        llm_response_content = response.choices[0].message.content.strip()

        # --- Conceptual Guardrail Integration Point (Output) ---
        # (response_postprocessor.py from doc 3.1 could live here)
        if st.session_state.rag_enabled and context_chunks: # Only apply context-based guards if RAG was used
            if st.session_state.guardrails_enabled.get("conli", False):
                is_conli_ok, conli_msg = conli_guard_validator(llm_response_content, context_chunks)
                st.sidebar.caption(f"Conli Guard: {conli_msg}")
                # if not is_conli_ok: llm_response_content += f"\n[Warning: {conli_msg}]"

            if st.session_state.guardrails_enabled.get("cove", False):
                is_cove_ok, cove_msg = cove_guard_validator(llm_response_content, context_chunks, query)
                st.sidebar.caption(f"CoVe Guard: {cove_msg}")
                # if not is_cove_ok: llm_response_content += f"\n[Warning: {cove_msg}]"
        
        # ContextCheck guard could look at chat history (not implemented here)
        # if st.session_state.guardrails_enabled.get("contextcheck", False):
        # is_context_ok, context_msg = contextcheck_guard_validator(llm_response_content, st.session_state.get('chat_history', []))
        # st.sidebar.caption(f"ContextCheck Guard: {context_msg}")

        if st.session_state.guardrails_enabled.get("detectpii", False):
            is_pii_ok_response, pii_guard_msg_response = detectpii_guard_validator(llm_response_content)
            st.sidebar.caption(f"PII Guard (Response): {pii_guard_msg_response}")
            # if not is_pii_ok_response: llm_response_content = "[PII Redacted by Guardrail]"

        return llm_response_content
    except AuthenticationError:
        st.error("OpenAI API Key is invalid or not authorized. Please check your key and ensure it has credit/is active.")
        return "AuthenticationError: Invalid API Key or insufficient credits."
    except RateLimitError:
        st.error("You have exceeded your OpenAI API quota or rate limit. Please check your usage or try again later.")
        return "RateLimitError: API quota exceeded."
    except APIError as e:
        st.error(f"An OpenAI API error occurred: {e}")
        return f"OpenAI API Error: {e}"
    except Exception as e:
        st.error(f"An unexpected error occurred: {str(e)}")
        return f"An unexpected error occurred: {str(e)}"

# --- File Extraction Helper Functions (Pipeline Step 2: Extraction) ---
def extract_text_from_pdf(file_bytes):
    if not py_pdf_available:
        st.error("PyPDF2 is not available. Cannot process PDF.")
        return None
    try:
        pdf_reader = PyPDF2.PdfReader(BytesIO(file_bytes))
        text = ""
        for page_num in range(len(pdf_reader.pages)):
            page = pdf_reader.pages[page_num]
            text += page.extract_text() or "" # Ensure to handle None from extract_text
        return text
    except Exception as e:
        st.error(f"Erreur lors de l'extraction du texte du PDF: {e}")
        return None

def extract_text_from_docx(file_bytes): # NEWLY ADDED as per Doc 2.1.1, 3.3
    if not python_docx_available:
        st.error("python-docx is not available. Cannot process DOCX.")
        return None
    try:
        doc = docx.Document(BytesIO(file_bytes))
        text = "\n".join([paragraph.text for paragraph in doc.paragraphs])
        return text
    except Exception as e:
        st.error(f"Erreur lors de l'extraction du texte du DOCX: {e}")
        return None

def extract_text_from_html(file_bytes):
    if not beautiful_soup_available:
        st.error("BeautifulSoup4 is not available. Cannot process HTML.")
        return None
    try:
        soup = BeautifulSoup(file_bytes, "html.parser")
        # Remove common non-content tags before text extraction
        for script_or_style in soup(["script", "style", "header", "footer", "nav", "aside"]):
            script_or_style.decompose()
        text = soup.get_text(separator=" ", strip=True)
        return text
    except Exception as e:
        st.error(f"Erreur lors de l'extraction du texte du HTML (fichier): {e}")
        return None

def extract_text_from_image_easyocr(file_bytes): # Doc 2.1.1 suggests Tesseract, code uses EasyOCR
    global easyocr_reader_en_fr # Use the global cached reader
    if not easy_ocr_available or not pillow_available:
        st.error("EasyOCR ou Pillow ne sont pas disponibles/configurés. Impossible de traiter l'image.")
        return None
    if easyocr_reader_en_fr is None:
        st.error("Le lecteur EasyOCR n'a pas pu être initialisé. Impossible de traiter l'image.")
        return None
    try:
        image_bytes_for_ocr = BytesIO(file_bytes)
        # Convert image to a format EasyOCR can handle (e.g., numpy array)
        image_np = np.array(Image.open(image_bytes_for_ocr)) # Requires Pillow
        result = easyocr_reader_en_fr.readtext(image_np, detail=0, paragraph=True)
        return "\n".join(result)
    except Exception as e:
        st.error(f"Erreur lors de l'OCR de l'image avec EasyOCR: {e}")
        return None

def fetch_and_extract_text_from_url(url):
    if not requests_available:
        st.error("Le module 'requests' n'est pas disponible. Impossible de récupérer le contenu Web.")
        return None, None
    if not beautiful_soup_available:
        st.error("Le module 'BeautifulSoup4' n'est pas disponible. Impossible d'analyser le contenu HTML.")
        return None, None
    try:
        headers = { # Mimic a browser to avoid simple blocks
            'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/91.0.4472.124 Safari/537.36',
            'Accept': 'text/html,application/xhtml+xml,application/xml;q=0.9,image/webp,*/*;q=0.8',
            'Accept-Language': 'en-US,en;q=0.5',
        }
        response = requests.get(url, headers=headers, timeout=15, allow_redirects=True)
        response.raise_for_status() # Raise HTTPError for bad responses (4XX or 5XX)
        
        doc_name = url # Default doc name to URL
        soup = BeautifulSoup(response.text, "html.parser")

        if soup.title and soup.title.string: # Try to get a better doc name from title
            doc_name = soup.title.string.strip()
        
        # Remove common non-content tags
        for non_content_tag in soup(["script", "style", "header", "footer", "nav", "aside", "form", "meta", "link"]):
            non_content_tag.decompose()
        
        text = soup.get_text(separator=" ", strip=True) # Get text, stripping tags and extra whitespace
        return text, doc_name
    except requests.exceptions.Timeout:
        st.error(f"Timeout lors de la récupération de l'URL '{url}'.")
        return None, None
    except requests.exceptions.RequestException as e:
        st.error(f"Erreur lors de la récupération de l'URL '{url}': {e}")
        return None, None
    except Exception as e: # Catch-all for other BS4 errors etc.
        st.error(f"Erreur lors de l'analyse du contenu HTML de '{url}': {e}")
        return None, None

# --- Streamlit App UI ---
st.title("Agent RAG - Étude de Faisabilité et Architecture")
st.markdown("Basé sur le document du 20 Juin 2025. Ingestion: PDF, DOCX, HTML, Images (EasyOCR), Liens Web. Stockage: FAISS.")

# Display missing dependencies warnings
if not py_pdf_available: st.warning("PyPDF2 non installé. Traitement PDF indisponible. `pip install PyPDF2`")
if not python_docx_available: st.warning("python-docx non installé. Traitement DOCX indisponible. `pip install python-docx`") # New
if not beautiful_soup_available: st.warning("BeautifulSoup4 non installé. Traitement HTML indisponible. `pip install beautifulsoup4`")
if not requests_available: st.warning("Requests non installé. URLs indisponibles. `pip install requests`")
if not pillow_available: st.warning("Pillow non installé. Traitement image (EasyOCR) indisponible. `pip install Pillow`")
if not easy_ocr_available or easyocr_reader_en_fr is None:
    st.warning("EasyOCR non installé/initialisé correctement. OCR indisponible. `pip install easyocr` (et potentiellement des dépendances système pour l'OCR).")
if not faiss_available: st.error("FAISS non installé. Le stockage et la recherche vectorielle sont indisponibles. `pip install faiss-cpu` ou `faiss-gpu`")

st.markdown("---")

# --- Sidebar Configuration ---
st.sidebar.title("Configuration (Conforme Doc 2.1.1)")
st.sidebar.markdown("**IMPORTANT:** Clé API OpenAI requise.")
api_key_input = st.sidebar.text_input(
    "Clé API OpenAI:", type="password", help="Obtenez sur platform.openai.com."
)
if api_key_input:
    st.session_state.openai_api_key = api_key_input
elif os.getenv("OPENAI_API_KEY") and not st.session_state.openai_api_key: # Check env var if not already set
     st.session_state.openai_api_key = os.getenv("OPENAI_API_KEY")

st.sidebar.selectbox(
    "Modèle d'Embedding OpenAI (Doc 2.1.1):",
    list(EMBEDDING_DIMS.keys()),
    key='embedding_model', 
    help="Modèle pour vectorisation. 'text-embedding-3-small' est recommandé par le document."
)
st.sidebar.info(f"Modèle d'embedding actif: `{st.session_state.embedding_model}`")

# --- Conceptual Guardrails Sidebar (Section 3.4) ---
st.sidebar.subheader("Validateurs Guardrails (Conceptuel - Doc 3.4)")
st.sidebar.caption("Activation/Désactivation (simulation).")
st.session_state.guardrails_enabled["conli"] = st.sidebar.checkbox("conli-guard (cohérence LLM/connaissances)", value=st.session_state.guardrails_enabled["conli"])
st.session_state.guardrails_enabled["cove"] = st.sidebar.checkbox("cove-guard (justification logique)", value=st.session_state.guardrails_enabled["cove"])
st.session_state.guardrails_enabled["contextcheck"] = st.sidebar.checkbox("contextcheck-guard (contexte conversationnel)", value=st.session_state.guardrails_enabled["contextcheck"])
st.session_state.guardrails_enabled["detectpii"] = st.sidebar.checkbox("detectpii-guard (détection PII)", value=st.session_state.guardrails_enabled["detectpii"])
st.session_state.guardrails_enabled["unusual_prompt"] = st.sidebar.checkbox("unusual-prompt-guard (requêtes suspectes)", value=st.session_state.guardrails_enabled["unusual_prompt"])
st.sidebar.markdown("---")
# --- End Guardrails Sidebar ---

st.header("1. Agent d'Acquisition de Documents (Indexation FAISS)")
st.markdown("Pipeline: Import -> Extraction -> Découpage -> Vectorisation -> Indexation FAISS (Doc 3.2)")

st.session_state.rag_enabled = st.checkbox(
    "Activer la Recherche Augmentée (RAG)",
    value=st.session_state.get('rag_enabled', True)
)
if st.session_state.rag_enabled:
    st.info("Mode RAG activé: L'exploitation utilisera les documents indexés dans FAISS.")
else:
    st.warning("Mode RAG désactivé: L'exploitation répondra sans consulter les documents (LLM seul).")


acquisition_tab1, acquisition_tab2, acquisition_tab3 = st.tabs(["Fichier", "Texte Brut", "URL Web"])
uploaded_file_val = None
doc_text_area_content_val = ""
doc_name_manual_val = ""
web_url_input_val = ""

with acquisition_tab1:
    st.markdown("Formats supportés (Doc 2.1.1): PDF, DOCX, HTML, TXT, PNG, JPG, JPEG.") # Added DOCX
    uploaded_file_val = st.file_uploader(
        "Choisissez un fichier document:",
        type=["pdf", "txt", "html", "png", "jpg", "jpeg", "docx"], # Added docx
        key=f"file_uploader_k{st.session_state.file_uploader_key_suffix}"
    )
    if uploaded_file_val:
        st.info(f"Fichier sélectionné: **{uploaded_file_val.name}**. Sera priorisé pour le traitement.")

with acquisition_tab2:
    doc_name_manual_val = st.text_input(
        "Nom du document (optionnel, si texte collé):",
        placeholder="Ex: Notes de réunion du 15/05",
        key=f"doc_name_manual_input_k{st.session_state.doc_name_manual_input_key_suffix}"
    )
    doc_text_area_content_val = st.text_area(
        "Ou collez le contenu du document ici:",
        height=200,
        key=f"doc_text_area_input_k{st.session_state.doc_text_area_input_key_suffix}"
    )
    if doc_text_area_content_val and not uploaded_file_val: # Prioritize file upload
        st.info("Texte collé détecté. Sera traité si aucun fichier n'est téléversé.")

with acquisition_tab3:
    web_url_input_val = st.text_input(
        "Entrez l'URL d'une page web à scraper:",
        placeholder="Ex: https://www.streamlit.io",
        key=f"web_url_input_k{st.session_state.web_url_input_key_suffix}"
    )
    if web_url_input_val and not uploaded_file_val: # Prioritize file upload
        st.info("URL détectée. Sera traitée si aucun fichier n'est téléversé et aucun texte collé n'est fourni.")

col1_acq, col2_acq = st.columns(2)
with col1_acq:
    chunk_method = st.selectbox("Méthode de Découpage (Chunking):", ["Paragraphe ('\\n\\n')", "Taille Fixe (Mots)"])
with col2_acq:
    if chunk_method == "Taille Fixe (Mots)":
        chunk_size_words = st.slider("Taille des Chunks (en mots):", 50, 500, 150, 10)
        chunk_overlap_words = st.slider("Chevauchement des Chunks (en mots):", 0, 100, 20, 5)

if st.button("Traiter et Vectoriser dans FAISS (Pipeline étapes 1-4)"):
    if not st.session_state.openai_api_key:
        st.error("Clé API OpenAI manquante. Veuillez la configurer dans la barre latérale.")
    elif not faiss_available:
        st.error("FAISS n'est pas disponible. Impossible de traiter et stocker les documents.")
    else:
        current_embedding_dim = EMBEDDING_DIMS.get(st.session_state.embedding_model)
        if not current_embedding_dim:
            st.error(f"Dimension d'embedding inconnue pour le modèle {st.session_state.embedding_model}.")
            st.stop()

        if st.session_state.faiss_index is None or st.session_state.faiss_index_dim != current_embedding_dim:
            if st.session_state.faiss_index is not None and st.session_state.faiss_index_dim != current_embedding_dim:
                st.warning(f"Le modèle d'embedding a changé (dim: {st.session_state.faiss_index_dim} -> {current_embedding_dim}). L'index FAISS et les chunks existants vont être effacés.")
            
            st.session_state.faiss_index = faiss.IndexFlatIP(current_embedding_dim) # Cosine similarity for normalized vectors
            st.session_state.faiss_index_dim = current_embedding_dim
            st.session_state.chunk_store = [] 
            st.info(f"Index FAISS (re)initialisé pour dimension {current_embedding_dim} (modèle: {st.session_state.embedding_model}).")

        doc_text_to_process = None
        doc_name_for_processing = "Document Inconnu"
        source_type_for_processing = "inconnu"
        input_source_used = None # To track which input needs reset

        if uploaded_file_val is not None:
            input_source_used = 'file'
            doc_name_for_processing = uploaded_file_val.name
            file_bytes = uploaded_file_val.getvalue()
            source_type_for_processing = uploaded_file_val.type
            with st.spinner(f"Étape 2: Extraction du texte de '{doc_name_for_processing}'..."):
                if uploaded_file_val.type == "application/pdf" and py_pdf_available:
                    doc_text_to_process = extract_text_from_pdf(file_bytes)
                elif uploaded_file_val.type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document" and python_docx_available: # MIME type for .docx
                    doc_text_to_process = extract_text_from_docx(file_bytes)
                elif uploaded_file_val.type == "text/html" and beautiful_soup_available:
                    doc_text_to_process = extract_text_from_html(file_bytes)
                elif uploaded_file_val.type == "text/plain":
                    try: doc_text_to_process = file_bytes.decode('utf-8')
                    except UnicodeDecodeError:
                        st.warning("Décodage UTF-8 échoué, tentative avec 'latin-1'.")
                        try: doc_text_to_process = file_bytes.decode('latin-1')
                        except UnicodeDecodeError: st.error("Impossible de décoder le fichier TXT."); doc_text_to_process = None
                elif uploaded_file_val.type in ["image/png", "image/jpeg"] and easy_ocr_available and pillow_available and easyocr_reader_en_fr:
                    doc_text_to_process = extract_text_from_image_easyocr(file_bytes)
                elif uploaded_file_val.type in ["image/png", "image/jpeg"]:
                    if not (easy_ocr_available and easyocr_reader_en_fr): st.error("EasyOCR non disponible ou non initialisé pour le traitement d'image.")
                    if not pillow_available: st.error("Pillow non disponible pour le traitement d'image.")
                    doc_text_to_process = None # Ensure it's None if prerequisites fail
                else:
                    st.warning(f"Type de fichier '{uploaded_file_val.type}' non supporté ou dépendance manquante."); doc_text_to_process = None
            
            if not doc_text_to_process or not doc_text_to_process.strip():
                 if uploaded_file_val: st.error(f"Extraction de texte échouée ou le document '{doc_name_for_processing}' est vide.")

        elif web_url_input_val.strip() and requests_available and beautiful_soup_available:
            input_source_used = 'url'
            with st.spinner(f"Étape 2: Récupération et extraction du contenu de l'URL: {web_url_input_val}..."):
                doc_text_to_process, scraped_doc_name = fetch_and_extract_text_from_url(web_url_input_val)
                if doc_text_to_process:
                    doc_name_for_processing = scraped_doc_name or web_url_input_val # Use scraped name if available
                    source_type_for_processing = "URL web"
                else: st.error(f"Extraction de texte échouée pour l'URL: '{web_url_input_val}'.")
        elif web_url_input_val.strip() and (not requests_available or not beautiful_soup_available):
            st.error("Les modules 'requests' ou 'BeautifulSoup4' sont manquants pour traiter les URLs.")

        elif doc_text_area_content_val.strip():
            input_source_used = 'text_area'
            doc_text_to_process = doc_text_area_content_val
            doc_name_for_processing = doc_name_manual_val.strip() or f"Texte Collé Auto-Nommé {st.session_state.doc_counter + 1}"
            source_type_for_processing = "texte collé"
        else:
            st.warning("Veuillez fournir un document via un fichier, une URL, ou en collant du texte.")

        processed_successfully = False
        if doc_text_to_process and doc_text_to_process.strip():
            with st.spinner(f"Traitement de '{doc_name_for_processing}': Pré-traitement, Découpage (Étape 3), Vectorisation et Indexation FAISS (Étape 4)..."):
                st.session_state.doc_counter += 1
                doc_id = f"doc_{st.session_state.doc_counter}_{str(uuid.uuid4())[:8]}" # Unique ID for the document
                
                # Store basic document info (not chunks yet)
                st.session_state.documents[doc_id] = {
                    'name': doc_name_for_processing, 'original_text_preview': doc_text_to_process[:500] + "...",
                    'source_type': source_type_for_processing, 'uploaded_at': datetime.now().isoformat()
                }

                cleaned_text = preprocess_text(doc_text_to_process) # query_preprocessor.py equivalent
                
                detected_lang = "N/A" # Default language
                try: # Language detection (Doc 3.2, 3.3)
                    if cleaned_text: # Ensure there's text to detect
                        detected_lang = detect(cleaned_text[:min(500, len(cleaned_text))]) # Detect on a sample
                except LangDetectException: st.warning(f"Détection de langue échouée pour '{doc_name_for_processing}'. Langue marquée N/A.")
                except Exception as e_lang: st.warning(f"Erreur inattendue pendant la détection de langue: {e_lang}")


                new_chunks_text = []
                if chunk_method == "Paragraphe ('\\n\\n')":
                    new_chunks_text = [p.strip() for p in cleaned_text.split('\n\n') if p.strip()]
                elif chunk_method == "Taille Fixe (Mots)":
                    new_chunks_text = simple_chunker(cleaned_text, chunk_size=chunk_size_words, overlap=chunk_overlap_words)

                doc_chunk_count = 0
                total_chunks_to_embed = len(new_chunks_text)

                if total_chunks_to_embed > 0 and st.session_state.faiss_index is not None:
                    prog_bar = st.progress(0, text=f"Vectorisation & Indexation FAISS: 0/{total_chunks_to_embed} chunks")
                    for i_chunk, chunk_text in enumerate(new_chunks_text):
                        if chunk_text: # Ensure chunk is not empty
                            embedding = get_embedding(chunk_text, st.session_state.openai_api_key, model=st.session_state.embedding_model)
                            if embedding:
                                embedding_np = np.array([embedding], dtype='float32')
                                try:
                                    st.session_state.faiss_index.add(embedding_np) # Add to FAISS
                                    # Store metadata corresponding to the FAISS index position
                                    chunk_metadata_item = {
                                        'id': str(uuid.uuid4()), 'doc_id': doc_id, 'doc_name': doc_name_for_processing,
                                        'chunk_index_in_doc': i_chunk, 'text': chunk_text,
                                        'metadata': { # Detailed metadata
                                            'source_doc_id': doc_id, 'source_doc_name': doc_name_for_processing,
                                            'source_type': source_type_for_processing, 'char_length': len(chunk_text),
                                            'word_count': len(chunk_text.split()), 'lang': detected_lang,
                                            'created_at': datetime.now().isoformat()}
                                    }
                                    st.session_state.chunk_store.append(chunk_metadata_item)
                                    doc_chunk_count +=1
                                except Exception as e_faiss:
                                    st.error(f"Erreur lors de l'ajout du chunk {i_chunk+1} à FAISS: {e_faiss}")
                                    # Decide if you want to break or continue
                            else:
                                st.error(f"Échec de la vectorisation du chunk {i_chunk+1} du document '{doc_name_for_processing}'. Chunk ignoré.")
                        prog_bar.progress((i_chunk + 1) / total_chunks_to_embed, text=f"Vectorisation & Indexation FAISS: {i_chunk+1}/{total_chunks_to_embed} chunks")
                    prog_bar.empty() # Remove progress bar after completion

                if doc_chunk_count > 0:
                    st.success(f"Document '{doc_name_for_processing}' (ID: {doc_id}): {doc_chunk_count} chunks ont été traités et ajoutés à l'index FAISS. Total dans FAISS: {st.session_state.faiss_index.ntotal}.")
                    processed_successfully = True
                elif not new_chunks_text and cleaned_text: # Text was present but no chunks generated
                    st.warning(f"Aucun chunk n'a été généré pour le document '{doc_name_for_processing}'. Vérifiez la méthode de découpage et le contenu.")
                elif not cleaned_text and doc_text_to_process: # Text became empty after preprocessing
                     st.warning(f"Le document '{doc_name_for_processing}' est devenu vide après le pré-traitement. Aucun chunk généré.")
        
        if processed_successfully and input_source_used: # Reset the input field that was used
            if input_source_used == 'file': st.session_state.file_uploader_key_suffix += 1
            elif input_source_used == 'url': st.session_state.web_url_input_key_suffix += 1
            elif input_source_used == 'text_area':
                st.session_state.doc_text_area_input_key_suffix += 1
                st.session_state.doc_name_manual_input_key_suffix += 1
            st.rerun() # Rerun to clear input fields and update UI

# --- Display Chunks ---
st.subheader("Aperçu des Derniers Chunks Indexés (Métadonnées de `chunk_store`)")
if faiss_available and st.session_state.chunk_store and st.session_state.faiss_index and st.session_state.faiss_index.ntotal > 0:
    num_to_display = min(5, len(st.session_state.chunk_store)) # Display last 5 or fewer
    # Prepare data for display, fetching from the end of chunk_store
    display_data = [{
        "ID Chunk": c['id'][:8], # Short ID for display
        "Nom Doc": c.get('doc_name', c.get('doc_id', 'N/A')), # Fallback to doc_id
        "Source": c['metadata'].get('source_type', 'N/A'),
        "Langue (Doc)": c['metadata'].get('lang', 'N/A'), # Language detected for the document
        "Aperçu Texte": c['text'][:70] + "..." if len(c['text']) > 70 else c['text'],
        "Indexé FAISS": "Oui" # Assumed if in chunk_store and FAISS is active
    } for c in reversed(st.session_state.chunk_store[-num_to_display:])] # Iterate reversed for latest first
    
    if display_data:
        st.dataframe(pd.DataFrame(display_data), use_container_width=True, hide_index=True)
    
    st.caption(f"Total de chunks dans l'index FAISS: {st.session_state.faiss_index.ntotal}. Métadonnées stockées pour: {len(st.session_state.chunk_store)} chunks.")
    if st.session_state.faiss_index.ntotal != len(st.session_state.chunk_store):
        st.warning("Incohérence détectée: Le nombre de chunks dans FAISS ne correspond pas au nombre de métadonnées stockées. Cela peut arriver si une erreur survient pendant l'indexation.")
elif not faiss_available:
     st.info("FAISS non disponible. L'affichage des chunks indexés est impossible.")
else:
    st.info("La base vectorielle FAISS est vide ou non initialisée. Veuillez traiter des documents pour l'alimenter.")
st.markdown("---")

# --- 2. Agent d'Exploitation ---
st.header("2. Agent d'Exploitation (Requête et Génération)")
exploit_msg = "Pipeline (Doc 3.2 étapes 5-8): Requête -> Vectorisation -> Recherche FAISS -> Génération GPT -> Affichage." if st.session_state.rag_enabled else "Mode RAG désactivé. L'agent répondra en utilisant uniquement ses connaissances générales (LLM seul)."
st.markdown(exploit_msg)

user_query = st.text_input("Posez votre question ici (Étape 5: Requête utilisateur):", key="user_query_input_main", placeholder="Ex: Quel est le rôle de FAISS dans ce projet ?")

if st.session_state.rag_enabled:
    col1_exp, col2_exp = st.columns(2)
    with col1_exp: top_k_retrieval = st.slider("Nombre de Chunks à récupérer (Top-K):", 1, 10, 3, help="Nombre de chunks les plus similaires à récupérer de FAISS.")
    with col2_exp: similarity_thresh = st.slider("Seuil de Similarité Minimum:", 0.0, 1.0, 0.70, 0.05, help="Similarité minimale pour qu'un chunk soit considéré pertinent.")

if st.button("Obtenir Réponse (OpenAI avec/sans RAG FAISS)"):
    if not st.session_state.openai_api_key: st.error("Clé API OpenAI manquante! Veuillez la configurer.")
    elif not user_query: st.warning("Veuillez entrer une question.")
    elif not faiss_available and st.session_state.rag_enabled:
        st.error("FAISS non disponible. La recherche RAG est impossible. Désactivez le RAG ou installez FAISS.")
    else:
        retrieved_chunks_for_llm = []; query_embedding_success = True # Assume success
        
        if not st.session_state.rag_enabled: 
            st.info("Mode RAG désactivé. Génération de réponse sans consultation des documents.")
        elif st.session_state.faiss_index is None or st.session_state.faiss_index.ntotal == 0 :
            st.warning("Base de données vectorielle FAISS vide ou non initialisée. Le mode RAG est actif mais aucun document n'est disponible pour la recherche.")
        else: # RAG enabled, FAISS available and has data
            with st.spinner("Étape 6: Vectorisation de la requête et recherche sémantique dans FAISS..."):
                query_preprocessed = preprocess_text(user_query)
                query_embedding = get_embedding(query_preprocessed, st.session_state.openai_api_key, model=st.session_state.embedding_model)
                if query_embedding:
                    retrieved_chunks_for_llm = faiss_similarity_search(
                        query_embedding, st.session_state.faiss_index, st.session_state.chunk_store,
                        top_k=top_k_retrieval, similarity_threshold=similarity_thresh
                    )
                    st.subheader(f"Chunks contextuels récupérés de FAISS ({len(retrieved_chunks_for_llm)}):")
                    if retrieved_chunks_for_llm:
                        for i, chunk_info in enumerate(retrieved_chunks_for_llm):
                            doc_display_name = chunk_info.get('doc_name', chunk_info.get('doc_id', 'Inconnu'))
                            similarity_score = chunk_info.get('score', 0.0)
                            lang_info = chunk_info.get('metadata',{}).get('lang','N/A')
                            with st.expander(f"Chunk Pertinent {i+1} (Doc: '{doc_display_name}', Similarité: {similarity_score:.4f})", expanded=(i==0)): # Expand first by default
                                st.markdown(f"**Langue (Doc):** `{lang_info}`")
                                st.markdown(f"**Texte du Chunk:**\n\n{chunk_info['text']}")
                    else: st.info("Aucun chunk pertinent n'a été trouvé dans FAISS avec les critères actuels. La réponse sera générée sans contexte documentaire.")
                else: 
                    st.error("Échec de la vectorisation de la requête. La réponse sera générée sans contexte documentaire.")
                    query_embedding_success = False # Mark as failed
        
        with st.spinner("Étape 7: Génération de la réponse via le LLM OpenAI..."):
            # Use gpt-4 or gpt-4o as per doc 2.1.1
            llm_model_to_use = "gpt-4o" # Default to gpt-4o as per latest in doc
            final_answer = real_llm_generation_with_openai(
                user_query, 
                retrieved_chunks_for_llm if query_embedding_success else [], # Pass empty if embedding failed
                st.session_state.openai_api_key,
                model=llm_model_to_use 
            )
            
            answer_header = f"Réponse (Modèle LLM: {llm_model_to_use}"
            if st.session_state.rag_enabled and query_embedding_success and retrieved_chunks_for_llm: 
                answer_header += " - avec contexte RAG/FAISS)"
            elif st.session_state.rag_enabled: 
                answer_header += " - RAG actif, mais sans contexte FAISS pertinent/disponible)"
            else: 
                answer_header += " - RAG désactivé, LLM seul)"
            
            st.subheader(answer_header)
            st.markdown(final_answer) # Étape 8: Affichage final
st.markdown("---")


# --- Sidebar System Info & Reset ---
st.sidebar.header("État du Système RAG")
st.sidebar.markdown(
    "- **Acquisition (Doc 3.1 `index_builder`):** Vectorise les documents, stocke dans FAISS et `chunk_store` (en mémoire).\n"
    "- **Exploitation (Doc 3.1 `rag_pipeline`):** Génère des réponses augmentées si RAG est actif."
)
st.sidebar.metric("Nombre de Documents Traités", st.session_state.doc_counter)
if faiss_available and st.session_state.faiss_index:
    st.sidebar.metric("Nombre de Chunks Indexés (FAISS)", st.session_state.faiss_index.ntotal)
    st.sidebar.metric("Nombre de Métadonnées de Chunks", len(st.session_state.chunk_store))
else:
    st.sidebar.metric("Chunks Indexés (FAISS)", 0)
    st.sidebar.metric("Métadonnées de Chunks", 0)

st.sidebar.markdown("---")
st.sidebar.markdown("**Extensions Futures (Doc 3.5):**")
st.sidebar.caption("- Intégration MongoDB (logs, supervision)")
st.sidebar.caption("- Authentification utilisateur")
# SharePoint/OneDrive (Doc 2.1.1) would also be a major future extension
st.sidebar.caption("- Connecteurs SharePoint / OneDrive (Doc 2.1.1)") 
st.sidebar.markdown("---")


if st.sidebar.button("Effacer Toutes les Données et Réinitialiser l'Agent"):
    st.session_state.documents = {}
    st.session_state.chunk_store = [] 
    st.session_state.doc_counter = 0
    if faiss_available: 
        st.session_state.faiss_index = None 
        st.session_state.faiss_index_dim = None

    # Reset input field keys to force re-render
    st.session_state.file_uploader_key_suffix += 1
    st.session_state.web_url_input_key_suffix += 1
    st.session_state.doc_name_manual_input_key_suffix += 1
    st.session_state.doc_text_area_input_key_suffix += 1
    
    # Reset Guardrails to default
    st.session_state.guardrails_enabled = {
        "conli": False, "cove": False, "contextcheck": False,
        "detectpii": False, "unusual_prompt": False
    }

    st.success("Toutes les données en mémoire (documents, chunks, index FAISS) et les champs d'entrée ont été réinitialisés.")
    st.rerun()